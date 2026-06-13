"""
XSIAM (Cortex XDR/XSIAM) API Client

Provides integration with Palo Alto Networks XSIAM using Advanced API key
authentication: the request is signed with a per-call nonce + timestamp,
hashed (SHA256) together with the API key, and sent in the Authorization
header.
"""

import gzip
import hashlib
import json
import logging
import re
import secrets
import sqlite3
import string
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
import requests

from my_config import get_config

logger = logging.getLogger(__name__)


class XsiamClient:
    """Client for interacting with the Palo Alto Cortex XSIAM/XDR API."""

    def __init__(self):
        self.config = get_config()
        self.api_key = self.config.xsiam_prod_api_key
        self.api_key_id = self.config.xsiam_prod_api_auth_id
        self.base_url = self.config.xsiam_prod_api_base_url
        self.ui_base_url = self._derive_ui_base_url()
        self.timeout = 60

        if not self.api_key:
            logger.warning("XSIAM API key not configured")
        if not self.api_key_id:
            logger.warning("XSIAM API key id not configured")
        if not self.base_url:
            logger.warning("XSIAM API base URL not configured")

    def is_configured(self) -> bool:
        return bool(self.api_key and self.api_key_id and self.base_url)

    def _derive_ui_base_url(self) -> Optional[str]:
        """UI base URL is the API base with the `api-` host prefix stripped.
        Override via XSIAM_PROD_UI_BASE_URL if the host pattern differs.
        """
        override = self.config.xsiam_prod_ui_base_url
        if override:
            return override.rstrip("/")
        if not self.base_url:
            return None
        return self.base_url.replace("//api-", "//", 1).rstrip("/")

    def case_url(self, case_id) -> Optional[str]:
        if not self.ui_base_url or case_id in (None, ""):
            return None
        return f"{self.ui_base_url}/case?caseId={case_id}"

    def _build_auth_headers(self) -> Dict[str, str]:
        """Build XSIAM advanced-auth headers (nonce + timestamp + SHA256 hash)."""
        nonce = "".join(
            secrets.choice(string.ascii_letters + string.digits) for _ in range(64)
        )
        timestamp = int(datetime.now(timezone.utc).timestamp()) * 1000
        auth_key = f"{self.api_key}{nonce}{timestamp}".encode("utf-8")
        api_key_hash = hashlib.sha256(auth_key).hexdigest()
        return {
            "x-xdr-timestamp": str(timestamp),
            "x-xdr-nonce": nonce,
            "x-xdr-auth-id": str(self.api_key_id),
            "Authorization": api_key_hash,
            "Content-Type": "application/json",
        }

    def _make_request(
        self,
        path: str,
        method: str = "POST",
        json_data: Optional[Dict[str, Any]] = None,
        params: Optional[Dict[str, Any]] = None,
    ) -> Dict[str, Any]:
        """Make an authenticated request to the XSIAM API.

        Args:
            path: API path (with or without leading slash).
            method: HTTP method (XSIAM endpoints are mostly POST).
            json_data: JSON body.
            params: Query string parameters.
        """
        if not self.is_configured():
            return {"error": "XSIAM API not configured (missing key, key id, or base URL)"}

        url = f"{self.base_url.rstrip('/')}/{path.lstrip('/')}"
        headers = self._build_auth_headers()

        try:
            logger.debug("XSIAM %s %s", method, path)
            response = requests.request(
                method=method,
                url=url,
                headers=headers,
                json=json_data if method != "GET" else None,
                params=params,
                timeout=self.timeout,
            )
            response.raise_for_status()
            if response.text:
                return response.json()
            return {"success": True}

        except requests.exceptions.HTTPError as e:
            status_code = e.response.status_code
            error_msg = e.response.text or str(e)
            if status_code == 401:
                return {"error": "Invalid XSIAM credentials or expired signature"}
            if status_code == 402:
                return {"error": "License does not cover this XSIAM API"}
            if status_code == 403:
                return {"error": "Access denied - insufficient XSIAM permissions"}
            if status_code == 404:
                return {"error": "XSIAM endpoint not found"}
            if status_code == 429:
                return {"error": "XSIAM API rate limit exceeded"}
            if status_code >= 500:
                return {"error": f"XSIAM server error: {status_code}"}
            logger.error("XSIAM API error: %s - %s", status_code, error_msg)
            return {"error": f"XSIAM API error ({status_code}): {error_msg}"}

        except requests.exceptions.Timeout:
            logger.error("XSIAM API request timed out")
            return {"error": "Request timed out"}

        except requests.exceptions.RequestException as e:
            logger.error("XSIAM request failed: %s", e)
            return {"error": f"Request failed: {e}"}

    # ==================== Public methods ====================

    def validate_auth(self) -> Dict[str, Any]:
        """Hit /api_keys/validate/ to confirm credentials and signing work."""
        return self._make_request("api_keys/validate/", method="POST", json_data={})

    @staticmethod
    def _build_request_data(
        filters: Optional[list] = None,
        search_from: int = 0,
        search_to: int = 100,
        sort: Optional[Dict[str, str]] = None,
    ) -> Dict[str, Any]:
        """Compose XSIAM's standard `request_data` envelope."""
        rd: Dict[str, Any] = {"search_from": search_from, "search_to": search_to}
        if filters:
            rd["filters"] = filters
        if sort:
            rd["sort"] = sort
        return rd

    @staticmethod
    def _time_filter(field: str, hours: int) -> Dict[str, Any]:
        """`field` >= now - hours, in epoch milliseconds."""
        cutoff_ms = int(
            (datetime.now(timezone.utc).timestamp() - hours * 3600) * 1000
        )
        return {"field": field, "operator": "gte", "value": cutoff_ms}

    def get_incidents(
        self,
        hours: Optional[int] = None,
        status: Optional[str] = None,
        limit: int = 50,
    ) -> Dict[str, Any]:
        """List incidents, optionally filtered by recency/status."""
        filters = []
        if hours and hours > 0:
            filters.append(self._time_filter("creation_time", hours))
        if status:
            filters.append({"field": "status", "operator": "in", "value": [status]})
        rd = self._build_request_data(
            filters=filters or None,
            search_to=max(1, min(limit, 100)),
            sort={"field": "creation_time", "keyword": "desc"},
        )
        return self._make_request(
            "public_api/v1/incidents/get_incidents/",
            method="POST",
            json_data={"request_data": rd},
        )

    def get_incident_extra_data(
        self, incident_id: str, alerts_limit: int = 50
    ) -> Dict[str, Any]:
        """Single incident with related alerts, artifacts, and network artifacts."""
        return self._make_request(
            "public_api/v1/incidents/get_incident_extra_data/",
            method="POST",
            json_data={
                "request_data": {
                    "incident_id": str(incident_id),
                    "alerts_limit": alerts_limit,
                }
            },
        )

    def update_incident(
        self,
        incident_id: str,
        status: Optional[str] = None,
        assigned_user_mail: Optional[str] = None,
        severity: Optional[str] = None,
        resolve_comment: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Update incident fields. Only provided fields are sent."""
        update_data: Dict[str, Any] = {}
        if status:
            update_data["status"] = status
        if assigned_user_mail:
            update_data["assigned_user_mail"] = assigned_user_mail
        if severity:
            update_data["severity"] = severity
        if resolve_comment:
            update_data["resolve_comment"] = resolve_comment
        return self._make_request(
            "public_api/v1/incidents/update_incident/",
            method="POST",
            json_data={
                "request_data": {
                    "incident_id": str(incident_id),
                    "update_data": update_data,
                }
            },
        )

    def get_alerts(
        self,
        hours: Optional[int] = None,
        severity: Optional[str] = None,
        limit: int = 100,
    ) -> Dict[str, Any]:
        """Query alerts via /public_api/v1/alerts/get_alerts_multi_events/."""
        filters = []
        if hours and hours > 0:
            filters.append(self._time_filter("creation_time", hours))
        if severity:
            filters.append(
                {"field": "severity", "operator": "in", "value": [severity]}
            )
        rd = self._build_request_data(
            filters=filters or None,
            search_to=max(1, min(limit, 100)),
            sort={"field": "creation_time", "keyword": "desc"},
        )
        return self._make_request(
            "public_api/v1/alerts/get_alerts_multi_events/",
            method="POST",
            json_data={"request_data": rd},
        )

    def get_endpoint(
        self,
        hostname: Optional[str] = None,
        ip: Optional[str] = None,
        endpoint_id: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Look up endpoints by hostname, IP, or endpoint_id."""
        filters = []
        if hostname:
            filters.append(
                {"field": "hostname", "operator": "in", "value": [hostname]}
            )
        if ip:
            filters.append({"field": "ip_list", "operator": "in", "value": [ip]})
        if endpoint_id:
            filters.append(
                {"field": "endpoint_id_list", "operator": "in", "value": [endpoint_id]}
            )
        if not filters:
            return {"error": "Provide at least one of hostname, ip, or endpoint_id"}
        return self._make_request(
            "public_api/v1/endpoints/get_endpoint/",
            method="POST",
            json_data={"request_data": {"filters": filters}},
        )

    # ==================== XQL ====================

    def start_xql_query(
        self,
        query: str,
        time_from_ms: Optional[int] = None,
        time_to_ms: Optional[int] = None,
        tenants: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """Submit an XQL query. Returns `{reply: "<query_id>"}` on success."""
        request_data: Dict[str, Any] = {
            "query": query,
            "tenants": tenants or [],
        }
        if time_from_ms is not None and time_to_ms is not None:
            request_data["_time"] = {"from": time_from_ms, "to": time_to_ms}
        return self._make_request(
            "public_api/v1/xql/start_xql_query/",
            method="POST",
            json_data={"request_data": request_data},
        )

    def get_query_results(
        self,
        query_id: str,
        poll: bool = True,
        poll_interval: float = 5.0,
        max_wait: float = 300.0,
    ) -> Dict[str, Any]:
        """Fetch XQL results for `query_id`. Polls while status is PENDING.

        Returns the full response dict; results live at
        `reply.results.data` (≤1000 rows inline) or `reply.results.stream_id`
        (use `get_query_results_stream` to fetch).
        """
        payload = {
            "request_data": {
                "query_id": query_id,
                "pending_flag": True,
                "format": "json",
            }
        }
        deadline = time.monotonic() + max_wait
        while True:
            res = self._make_request(
                "public_api/v1/xql/get_query_results/",
                method="POST",
                json_data=payload,
            )
            if "error" in res:
                return res
            status = (res.get("reply") or {}).get("status")
            if not poll or status != "PENDING":
                return res
            if time.monotonic() >= deadline:
                return {"error": f"XQL query {query_id} still PENDING after {max_wait}s"}
            time.sleep(poll_interval)

    def get_query_results_stream(self, stream_id: str) -> Dict[str, Any]:
        """Fetch the gzipped result stream for a large XQL result set.

        XSIAM returns NDJSON (one JSON object per line) gzipped in the body —
        not a JSON envelope — so this can't go through `_make_request`.
        """
        if not self.is_configured():
            return {"error": "XSIAM API not configured (missing key, key id, or base URL)"}

        url = f"{self.base_url.rstrip('/')}/public_api/v1/xql/get_query_results_stream/"
        headers = self._build_auth_headers()
        payload = {
            "request_data": {
                "stream_id": stream_id,
                "is_gzip_compressed": True,
            }
        }
        try:
            response = requests.post(
                url=url,
                headers=headers,
                json=payload,
                timeout=self.timeout,
            )
            response.raise_for_status()
            text = gzip.decompress(response.content).decode("utf-8")
            rows: List[Dict[str, Any]] = []
            for line in text.splitlines():
                line = line.strip()
                if not line:
                    continue
                try:
                    rows.append(json.loads(line))
                except json.JSONDecodeError:
                    logger.warning("XQL stream: skipping non-JSON line")
            return {"data": rows}
        except requests.exceptions.HTTPError as e:
            status_code = e.response.status_code
            return {"error": f"XSIAM stream error ({status_code}): {e.response.text or e}"}
        except (OSError, EOFError) as e:
            logger.error("XQL stream gzip decode failed: %s", e)
            return {"error": f"Stream decode failed: {e}"}
        except requests.exceptions.RequestException as e:
            logger.error("XSIAM stream request failed: %s", e)
            return {"error": f"Request failed: {e}"}

    # ==================== Correlation-rule deployment ====================

    # XSIAM severity enum values used by the correlation-rule "severity" field.
    _SEVERITY_ENUM = {
        "informational": "SEV_010_INFO",
        "low": "SEV_020_LOW",
        "medium": "SEV_030_MEDIUM",
        "high": "SEV_040_HIGH",
        "critical": "SEV_050_CRITICAL",
    }

    # Public-API path for creating a correlation rule. The correlation-rule
    # management surface is tenant/version dependent; this targets the
    # documented insert path. It is exercised ONLY on an explicit live deploy
    # (dry_run=False) — every other path validates the XQL and stops short of a
    # write, so a wrong path can never fire by accident. Override per call if
    # the tenant differs.
    CORRELATION_INSERT_PATH = "public_api/v1/correlations/insert/"

    def validate_xql(
        self, xql: str, window_hours: int = 1, max_wait: float = 60.0
    ) -> Dict[str, Any]:
        """Submit `xql` read-only over a short window purely to confirm it parses.

        Returns {"ok": True, "results": <n>} on success, {"error": ...} otherwise.
        This is the same engine the workbench dry-run uses — a query that runs
        here will run as a correlation rule.
        """
        xql = (xql or "").strip()
        if not xql:
            return {"error": "Empty XQL"}
        now_ms = int(time.time() * 1000)
        started = self.start_xql_query(xql, now_ms - window_hours * 3600 * 1000, now_ms)
        if "error" in started:
            return {"error": started["error"]}
        query_id = started.get("reply")
        if not query_id:
            return {"error": f"No query id returned ({started})"}
        res = self.get_query_results(query_id, poll=True, poll_interval=3.0, max_wait=max_wait)
        if "error" in res:
            return {"error": res["error"]}
        reply = res.get("reply") or {}
        return {"ok": True, "results": reply.get("number_of_results", 0)}

    def create_correlation_rule(
        self,
        name: str,
        xql: str,
        description: str = "",
        severity: str = "medium",
        mitre_tactics: Optional[List[str]] = None,
        mitre_techniques: Optional[List[str]] = None,
        search_window: str = "24_HOURS",
        enabled: bool = True,
        validate: bool = True,
        dry_run: bool = True,
        insert_path: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Create — or, by default, just preview — an XSIAM correlation rule.

        SAFE BY DEFAULT. With dry_run=True (the default) this validates the XQL
        against the live tenant read-only and returns exactly what WOULD be
        created, writing nothing. Pass dry_run=False to perform the live create;
        that is the single write path and is meant to be reached only behind an
        explicit, human-triggered deploy step.

        Returns one of:
          {"dry_run": True, "would_create": {...}, "validation": {...}}
          {"ok": True, "created": True, "reply": ...}
          {"error": "..."}
        """
        if not self.is_configured():
            return {"error": "XSIAM API not configured (missing key, key id, or base URL)"}
        name = (name or "").strip()
        xql = (xql or "").strip()
        if not name:
            return {"error": "Correlation rule needs a name"}
        if not xql:
            return {"error": "Correlation rule needs an XQL query"}

        validation = None
        if validate:
            validation = self.validate_xql(xql)
            if "error" in validation:
                return {"error": f"XQL did not validate: {validation['error']}",
                        "validated": False}

        rule = {
            "name": name,
            "description": description or name,
            "severity": self._SEVERITY_ENUM.get(severity.lower(), "SEV_030_MEDIUM"),
            "xql_query": xql,
            "search_window": search_window,
            "mitre_tactics": mitre_tactics or [],
            "mitre_techniques": mitre_techniques or [],
            "enabled": bool(enabled),
        }

        if dry_run:
            return {"dry_run": True, "validated": bool(validate),
                    "validation": validation, "would_create": rule}

        path = insert_path or self.CORRELATION_INSERT_PATH
        res = self._make_request(path, method="POST", json_data={"request_data": rule})
        if "error" in res:
            return {"error": res["error"], "created": False, "validation": validation}
        return {"ok": True, "created": True, "validation": validation,
                "reply": res.get("reply", res)}


# =========================================================================== #
# Vulnerability-export ingestion + reconciliation
#
# XSIAM exports a spreadsheet of vulnerability findings (one row per CVE x
# asset) to data/transient/. This block normalizes those rows into a local
# SQLite store and reconciles XSIAM's CVE coverage against our own CVE-triage
# results, so the triage app can act as a vendor-independent "backup" that
# cross-checks the commercial scanners (surface CVEs XSIAM flagged that we
# have not triaged).
#
# Entry points:
#     latest_export() -> Path | None
#     ingest_export(path=None) -> int
#     xsiam_cves() -> set[str]
#     enrich_cve(cve_id) -> dict | None
#     reconcile(triage_db_path=...) -> dict
# =========================================================================== #

# Both worktrees share the transient export drop; scan each.
_REPO_ROOT = Path(__file__).resolve().parent.parent
_TRANSIENT_DIRS = [
    _REPO_ROOT / "data" / "transient",
    Path("/home/vinay/security-ops-platform/data/transient"),
    Path("/home/vinay/security-ops-platform-dev/data/transient"),
]
_EXPORT_GLOB = "XSIAM_Vulnerability_Issues_*.xlsx"

VULN_DB_PATH = _REPO_ROOT / "data" / "transient" / "xsiam_findings.db"
DEFAULT_TRIAGE_DB = "/home/vinay/security-ops-platform-dev/data/transient/cve_triage_results.db"

CVE_RE = re.compile(r"CVE-\d{4}-\d{4,7}", re.IGNORECASE)

_COLUMNS = [
    "cve_id", "asset_name", "asset_type", "os", "os_distro",
    "affected_software", "package_version", "file_path", "package_in_use",
    "internet_exposed", "base_image_vuln", "exploitable", "exploit_level",
    "has_kev", "cvss_score", "cvss_severity", "epss", "fix_available",
    "fix_versions", "status", "location", "providers", "finding_sources",
    "source_file", "ingested_at",
]


# --------------------------- value parsing helpers ------------------------- #
def _yesno(val) -> Optional[int]:
    """Map a Yes/No-ish cell to 1/0; None when unknown."""
    if val is None:
        return None
    s = str(val).strip().lower()
    if s in ("yes", "true", "1", "y"):
        return 1
    if s in ("no", "false", "0", "n"):
        return 0
    return None


def _num(val) -> Optional[float]:
    """Parse a numeric cell that may arrive as int, float, or string."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        m = re.search(r"-?\d+(?:\.\d+)?", s)
        return float(m.group()) if m else None


def _text(val) -> Optional[str]:
    """Normalize a cell to a clean string, treating empties / 'None' as null."""
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() == "none":
        return None
    return s


def _split_cves(val) -> List[str]:
    """Extract every CVE ID from a cell that may list several."""
    if val is None:
        return []
    seen, out = set(), []
    for c in CVE_RE.findall(str(val)):
        cu = c.upper()
        if cu not in seen:
            seen.add(cu)
            out.append(cu)
    return out


# ------------------------------ export discovery --------------------------- #
def latest_export() -> Optional[Path]:
    """Newest XSIAM_Vulnerability_Issues_*.xlsx across both worktrees."""
    candidates: Dict[Path, Path] = {}
    for d in _TRANSIENT_DIRS:
        try:
            if not d.is_dir():
                continue
            for p in d.glob(_EXPORT_GLOB):
                candidates[p.resolve()] = p
        except OSError as e:
            logger.debug("scan failed for %s: %s", d, e)
    if not candidates:
        logger.warning("no XSIAM export found in %s", _TRANSIENT_DIRS)
        return None
    return max(candidates.values(), key=lambda p: p.stat().st_mtime)


# --------------------------------- storage --------------------------------- #
def _connect() -> sqlite3.Connection:
    VULN_DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(VULN_DB_PATH))
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS xsiam_findings (
            cve_id            TEXT,
            asset_name        TEXT,
            asset_type        TEXT,
            os                TEXT,
            os_distro         TEXT,
            affected_software TEXT,
            package_version   TEXT,
            file_path         TEXT,
            package_in_use    INT,
            internet_exposed  INT,
            base_image_vuln   INT,
            exploitable       INT,
            exploit_level     TEXT,
            has_kev           INT,
            cvss_score        REAL,
            cvss_severity     TEXT,
            epss              REAL,
            fix_available     INT,
            fix_versions      TEXT,
            status            TEXT,
            location          TEXT,
            providers         TEXT,
            finding_sources   TEXT,
            source_file       TEXT,
            ingested_at       TEXT
        )
        """
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_xsiam_findings_cve ON xsiam_findings(cve_id)"
    )
    return conn


def ingest_export(path: Optional[Path] = None) -> int:
    """Parse an XSIAM vulnerability export and store normalized findings.

    One DB row per (CVE x asset). When a cell lists multiple CVEs, one row is
    emitted per CVE. Idempotent: prior rows for the same source_file are
    deleted before insert, so re-ingesting the same file is safe.

    Returns the number of finding rows written.
    """
    if path is None:
        path = latest_export()
    if path is None:
        logger.error("ingest_export: no export available")
        return 0
    path = Path(path)
    if not path.exists():
        logger.error("ingest_export: file does not exist: %s", path)
        return 0

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        logger.error("ingest_export: failed to open %s: %s", path, e)
        return 0

    try:
        ws = wb.worksheets[0]  # active/first sheet holds the data
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            logger.warning("ingest_export: empty sheet in %s", path)
            return 0
        col = {h: i for i, h in enumerate(header) if h is not None}

        def cell(row, name):
            i = col.get(name)
            return row[i] if i is not None and i < len(row) else None

        source_file = path.name
        ingested_at = time.strftime("%Y-%m-%dT%H:%M:%S")
        records = []
        for row in rows:
            if row is None or all(v is None for v in row):
                continue
            cves = _split_cves(cell(row, "CVE ID"))
            if not cves:
                continue
            base = {
                "asset_name": _text(cell(row, "Asset Names")),
                "asset_type": _text(cell(row, "Asset Types")),
                "os": _text(cell(row, "Operating System")),
                "os_distro": _text(cell(row, "Operating System Distribution")),
                "affected_software": _text(cell(row, "Affected Software")),
                "package_version": _text(cell(row, "Software Package Version")),
                "file_path": _text(cell(row, "File Path")),
                "package_in_use": _yesno(cell(row, "Package in Use")),
                "internet_exposed": _yesno(cell(row, "Internet Exposed")),
                "base_image_vuln": _yesno(cell(row, "Base Image Vulnerability")),
                "exploitable": _yesno(cell(row, "Exploitable")),
                "exploit_level": _text(cell(row, "Exploit Level")),
                "has_kev": _yesno(cell(row, "Has KEV")),
                "cvss_score": _num(cell(row, "CVSS Score")),
                "cvss_severity": _text(cell(row, "CVSS Severity")),
                "epss": _num(cell(row, "EPSS Score")),
                "fix_available": _yesno(cell(row, "Fix Available")),
                "fix_versions": _text(cell(row, "Fix Versions")),
                "status": _text(cell(row, "Status")),
                "location": _text(cell(row, "Location")),
                "providers": _text(cell(row, "Providers")),
                "finding_sources": _text(cell(row, "Finding Sources")),
                "source_file": source_file,
                "ingested_at": ingested_at,
            }
            for cve in cves:
                rec = {"cve_id": cve, **base}
                records.append(tuple(rec[c] for c in _COLUMNS))
    finally:
        wb.close()

    conn = _connect()
    try:
        conn.execute(
            "DELETE FROM xsiam_findings WHERE source_file = ?", (source_file,)
        )
        placeholders = ", ".join("?" * len(_COLUMNS))
        conn.executemany(
            f"INSERT INTO xsiam_findings ({', '.join(_COLUMNS)}) "
            f"VALUES ({placeholders})",
            records,
        )
        conn.commit()
    finally:
        conn.close()

    logger.info("ingested %d findings from %s", len(records), source_file)
    return len(records)


# ----------------------------- query / aggregate --------------------------- #
def xsiam_cves() -> set:
    """Distinct CVE IDs ingested so far."""
    if not VULN_DB_PATH.exists():
        return set()
    conn = sqlite3.connect(str(VULN_DB_PATH))
    try:
        return {
            r[0] for r in conn.execute("SELECT DISTINCT cve_id FROM xsiam_findings")
        }
    except sqlite3.OperationalError:
        return set()
    finally:
        conn.close()


def enrich_cve(cve_id: str) -> Optional[dict]:
    """Aggregate all asset findings for one CVE into a single summary dict."""
    if not VULN_DB_PATH.exists():
        return None
    cve_id = cve_id.upper()
    conn = sqlite3.connect(str(VULN_DB_PATH))
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            "SELECT * FROM xsiam_findings WHERE cve_id = ?", (cve_id,)
        ).fetchall()
    except sqlite3.OperationalError:
        return None
    finally:
        conn.close()
    if not rows:
        return None

    def _max(field):
        vals = [r[field] for r in rows if r[field] is not None]
        return max(vals) if vals else None

    fix_versions = sorted({
        v.strip()
        for r in rows if r["fix_versions"]
        for v in str(r["fix_versions"]).split(",")
        if v.strip()
    })
    statuses = sorted({r["status"] for r in rows if r["status"]})

    return {
        "cve_id": cve_id,
        "asset_count": len(rows),
        "any_internet_exposed": bool(_max("internet_exposed")),
        "any_package_in_use": bool(_max("package_in_use")),
        "any_base_image_vuln": bool(_max("base_image_vuln")),
        "any_exploitable": bool(_max("exploitable")),
        "has_kev": bool(_max("has_kev")),
        "max_epss": _max("epss"),
        "max_cvss": _max("cvss_score"),
        "fix_available": bool(_max("fix_available")),
        "fix_versions": fix_versions,
        "statuses": statuses,
    }


# ---------------------- reconciliation vs our triage ----------------------- #
def reconcile(triage_db_path: str = DEFAULT_TRIAGE_DB) -> dict:
    """Compare XSIAM's CVE set against our triaged CVE set.

    xsiam_only = CVEs XSIAM flagged that aren't in our triage set; these are
    candidate gaps in our coverage. Guards gracefully if the triage DB is
    absent.
    """
    xsiam = xsiam_cves()

    our: set = set()
    triage_available = False
    p = Path(triage_db_path)
    if p.exists():
        conn = sqlite3.connect(str(p))
        try:
            our = {
                str(r[0]).upper()
                for r in conn.execute("SELECT cve_id FROM triage")
                if r[0]
            }
            triage_available = True
        except sqlite3.OperationalError as e:
            logger.warning("reconcile: triage DB unreadable: %s", e)
        finally:
            conn.close()
    else:
        logger.warning("reconcile: triage DB absent at %s", triage_db_path)

    return {
        "triage_available": triage_available,
        "overlap": sorted(xsiam & our),
        "xsiam_only": sorted(xsiam - our),
        "ours_only_count": len(our - xsiam),
        "xsiam_cve_count": len(xsiam),
        "our_cve_count": len(our),
    }


def _vuln_row_count() -> int:
    if not VULN_DB_PATH.exists():
        return 0
    conn = sqlite3.connect(str(VULN_DB_PATH))
    try:
        return conn.execute("SELECT COUNT(*) FROM xsiam_findings").fetchone()[0]
    except sqlite3.OperationalError:
        return 0
    finally:
        conn.close()


def _main(argv=None) -> int:
    import argparse

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    ap = argparse.ArgumentParser(
        description="XSIAM vulnerability-export ingest + reconcile"
    )
    ap.add_argument(
        "--ingest", nargs="?", const=True, default=None, metavar="PATH",
        help="ingest the given export (or the latest if PATH omitted) and print the row count",
    )
    args = ap.parse_args(argv)

    if args.ingest is not None:
        target = None if args.ingest is True else Path(args.ingest)
        n = ingest_export(target)
        print(f"ingested {n} findings")
        return 0

    latest = latest_export()
    cves = xsiam_cves()
    rec = reconcile()
    print("XSIAM vulnerability export summary")
    print(f"  latest export : {latest}")
    print(f"  findings rows : {_vuln_row_count()}")
    print(f"  distinct CVEs : {len(cves)}")
    print("  reconcile vs triage:")
    print(f"    triage DB available : {rec['triage_available']}")
    print(f"    xsiam CVE count     : {rec['xsiam_cve_count']}")
    print(f"    our CVE count       : {rec['our_cve_count']}")
    print(f"    overlap             : {len(rec['overlap'])}")
    print(f"    xsiam_only          : {len(rec['xsiam_only'])}")
    print(f"    ours_only           : {rec['ours_only_count']}")
    if rec["xsiam_only"]:
        print(f"    xsiam_only CVEs     : {', '.join(rec['xsiam_only'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(_main())
